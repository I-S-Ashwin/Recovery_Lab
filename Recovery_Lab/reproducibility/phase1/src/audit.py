"""Reproducible Phase 1 audit. Reads the source; never modifies it or trains models."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import platform
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl


# Original header, canonical field, type, unit, definition, role.
FIELDS = [
    ('Agmt Id', 'agreement_id', 'string', 'identifier', 'Agreement identifier; not a customer or vehicle identifier.', 'identifier'),
    ('Cust Age', 'customer_age', 'number', 'years', 'Recorded customer age; observation date is not supplied.', 'customer'),
    ('Cust Gender', 'customer_gender', 'string', 'code', 'Recorded gender category; audit only.', 'audit_only'),
    ('Cust Cibil Score', 'customer_cibil_score_raw', 'number', 'score/code', 'Source score preserved, including unresolved -1 code.', 'customer'),
    ('Cust Employment Type', 'customer_employment_code', 'string', 'code', 'Recorded employment code; codebook not provided.', 'customer'),
    ('Cust Net Salary', 'customer_net_salary_raw', 'number', 'INR; frequency unconfirmed', 'Recorded income amount; zero meaning and frequency unconfirmed.', 'customer'),
    ('Coborrower Flag', 'coborrower_flag', 'string', 'code', 'Recorded co-borrower flag; Y/N retained.', 'customer'),
    ('App Score Risk', 'application_risk_band', 'string', 'category', 'Application risk label; not a residual-risk outcome.', 'customer'),
    ('Agmt Date', 'agreement_date', 'date', 'date', 'Agreement event date.', 'event_date'),
    ('Seizure Date', 'seizure_date', 'date', 'date', 'Seizure event date; prediction-time anchor for recovery mode.', 'event_date'),
    ('Sold Date', 'sold_date', 'date', 'date', 'Sale event date; earliest assumed label availability, not system receipt time.', 'future'),
    ('Tenure', 'tenure_months', 'number', 'months; confirm', 'Recorded loan tenure; contractual interpretation needs confirmation.', 'policy'),
    ('Cust Net IRR', 'customer_net_irr_raw', 'number', 'reported percent; basis unconfirmed', 'Reported net interest measure; do not treat automatically as nominal annual rate.', 'policy'),
    ('Sold Date Month', 'sold_month_raw', 'number', 'month number', 'Redundant source sale month; reconcile with sold date.', 'future'),
    ('Sold Date Year', 'sold_year_raw', 'number', 'year', 'Redundant source sale year; reconcile with sold date.', 'future'),
    ('Seizure Date Month', 'seizure_month_raw', 'number', 'month number', 'Redundant source seizure month; reconcile with seizure date.', 'event_date'),
    ('Seizure Date Year', 'seizure_year_raw', 'number', 'year', 'Redundant source seizure year; reconcile with seizure date.', 'event_date'),
    ('Cust Branch', 'customer_branch', 'string', 'category', 'Recorded branch; historical assignment timestamp unavailable.', 'base_candidate'),
    ('Cust Region', 'customer_region', 'string', 'category', 'Business region code; not assumed to be a state.', 'base_candidate'),
    ('Cust State', 'customer_state', 'string', 'code', 'Recorded state code preserved without inferred remapping.', 'base_candidate'),
    ('Pincode Tier', 'pincode_tier', 'string', 'category', 'Provided settlement tier; full label retained.', 'base_candidate'),
    ('RC Availability', 'rc_availability', 'string', 'code', 'Registration certificate availability; observation time unverified.', 'conditional'),
    ('Registration Flag', 'registration_flag', 'string', 'code', 'Registration status; observation time unverified.', 'conditional'),
    ('Asset Disc Flag', 'asset_disc_flag', 'string', 'code', 'Binary source disc flag; exact business meaning unconfirmed.', 'conditional'),
    ('Asset Alloy Flag', 'asset_alloy_flag', 'string', 'code', 'Binary source alloy flag; exact business meaning unconfirmed.', 'conditional'),
    ('Asset Variant', 'asset_variant', 'string', 'category', 'Recorded vehicle variant.', 'base_candidate'),
    ('Asset Model', 'asset_model', 'string', 'category', 'Recorded vehicle model/segment label; levels are not assumed uniform.', 'base_candidate'),
    ('Asset Fuel Type', 'asset_fuel_type', 'string', 'category', 'Recorded propulsion category, EV/Petrol.', 'base_candidate'),
    ('Asset Cost At Disbursal', 'asset_cost_at_disbursal', 'number', 'INR', 'Recorded original asset cost.', 'base_candidate'),
    ('Loan Amount', 'loan_amount', 'number', 'INR', 'Original financed amount; policy input, not core physical valuation predictor.', 'policy'),
    ('LTV', 'ltv', 'number', 'fraction', 'Loan amount divided by original asset cost; source value preserved.', 'policy'),
    ('OS Balance At Liquidation', 'outstanding_balance_at_liquidation', 'number', 'INR', 'Recorded liquidation balance; unavailable at earlier decisions.', 'future'),
    ('Asset Age Months At Seizure', 'asset_age_months_at_seizure_raw', 'number', '30-day months', 'Source age label; matches agreement-to-seizure elapsed time in this file; manufacture age unverified.', 'age_unconfirmed'),
    ('Months Spent In Yard', 'months_spent_in_yard_raw', 'number', '30-day months', 'Realized seizure-to-sale delay; future information at seizure.', 'future'),
    ('Asset Bodycondition', 'asset_body_condition', 'string', 'code', 'Recorded body-condition code; mapping and inspection timestamp unconfirmed.', 'conditional'),
    ('Asset Tyrecondition', 'asset_tyre_condition', 'string', 'code', 'Recorded tyre-condition code; mapping and inspection timestamp unconfirmed.', 'conditional'),
    ('Asset Generalcondition', 'asset_general_condition', 'string', 'code', 'Recorded general-condition code; mapping and inspection timestamp unconfirmed.', 'conditional'),
    ('Asset Enginecondition', 'asset_engine_condition', 'string', 'code', 'Recorded engine-condition code; mapping and inspection timestamp unconfirmed.', 'conditional'),
    ('Asset Accident Flag', 'asset_accident_flag', 'string', 'code', 'Recorded accident flag; timing and history scope unconfirmed.', 'conditional'),
    ('Traiffic Challan Amount', 'traffic_challan_amount', 'number', 'INR', 'Recorded traffic fine amount; original misspelling mapped, timing unconfirmed.', 'conditional'),
    ('Target Sold Amount At Liquidation', 'sale_amount', 'number', 'INR', 'Observed gross liquidation sale amount; target only, costs not supplied.', 'target'),
]

DERIVED = [
    ('source_excel_row', 'integer', 'row number', 'Original worksheet row for traceability.', 'identifier'),
    ('customer_cibil_score_usable', 'number/null', 'score', 'Source score if nonnegative, otherwise null; no claim that -1 means no history.', 'customer'),
    ('cibil_special_code_flag', 'boolean', 'flag', 'True when source CIBIL is negative; original value retained.', 'customer'),
    ('customer_income_positive', 'number/null', 'INR; frequency unconfirmed', 'Source income only if positive; not verified monthly income.', 'customer'),
    ('income_nonpositive_flag', 'boolean', 'flag', 'True when recorded income is zero or negative.', 'customer'),
    ('months_since_agreement_at_seizure', 'number', '30-day months', '(Seizure date - agreement date) in days divided by 30; not confirmed vehicle age.', 'seizure_derived'),
    ('realized_yard_days', 'integer', 'days', '(Sold date - seizure date) in days; retrospective audit only.', 'future'),
    ('ltv_recomputed', 'number', 'fraction', 'Loan amount / asset cost; reconciliation and policy use only.', 'policy'),
]


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def scalar(x):
    if x is None or pd.isna(x):
        return None
    if isinstance(x, (datetime, pd.Timestamp)):
        return x.isoformat()
    if isinstance(x, np.generic):
        return x.item()
    return x


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + '\n', encoding='utf-8')


def write_jsonl(path, frame):
    with path.open('w', encoding='utf-8', newline='\n') as f:
        for row in frame.to_dict(orient='records'):
            f.write(json.dumps({k: scalar(v) for k, v in row.items()}, ensure_ascii=False, allow_nan=False) + '\n')


def availability(role):
    if role == 'base_candidate':
        return ('candidate; source timestamp not supplied', 'candidate; historical snapshot must be verified')
    if role == 'conditional':
        return ('exclude until meaning/availability confirmed', 'exclude future versions; require approval-time snapshot')
    if role == 'seizure_derived':
        return ('candidate at seizure; elapsed agreement time only', 'excluded; future event')
    if role == 'event_date':
        return ('event anchor/calendar only if known', 'agreement date only; future seizure fields excluded')
    if role == 'customer':
        return ('excluded from core physical valuation', 'policy context only after units/timing confirmation')
    if role == 'policy':
        return ('policy/audit input only', 'policy input only; not core valuation feature')
    if role == 'age_unconfirmed':
        return ('exclude raw duplicate; use explicitly named elapsed-agreement duration', 'excluded; future event')
    return ('excluded from predictors: ' + role, 'excluded from predictors: ' + role)


def audit(source: Path, out: Path):
    source = source.resolve()
    out = out.resolve()
    if source == out or source in out.parents:
        raise ValueError('Output must be separate from source file')
    before = sha(source)
    for folder in ['data', 'reports', 'configs']:
        (out / folder).mkdir(parents=True, exist_ok=True)
    book = openpyxl.load_workbook(source, read_only=True, data_only=False)
    if book.sheetnames != ['Sheet1']:
        raise ValueError(f'Unexpected sheets: {book.sheetnames}; explicitly review the schema')
    sheet = book['Sheet1']
    values = list(sheet.values)
    if list(values[0]) != [f[0] for f in FIELDS]:
        raise ValueError('Source schema differs from the documented 41-column contract')
    if any(isinstance(v, str) and v.startswith('=') for row in values[1:] for v in row):
        raise ValueError('Formula cells require a separate cached-value audit; refusing silent interpretation')
    indexed = [(i, row) for i, row in enumerate(values[1:], 2) if any(v is not None for v in row)]
    raw = pd.DataFrame([row for _, row in indexed], columns=values[0])
    excel_rows = [i for i, _ in indexed]
    d = raw.rename(columns={f[0]: f[1] for f in FIELDS}).copy()
    d.insert(0, 'source_excel_row', excel_rows)
    issues = []
    checks = []

    def check(name, mask, severity, explanation, threshold=None):
        mask = pd.Series(mask, index=d.index).fillna(False).astype(bool)
        checks.append({'rule': name, 'severity': severity, 'affected_rows': int(mask.sum()),
                       'explanation': explanation, 'threshold': threshold})
        for idx in d.index[mask]:
            issues.append({'source_excel_row': int(d.at[idx, 'source_excel_row']),
                           'agreement_id': scalar(d.at[idx, 'agreement_id']),
                           'rule': name, 'severity': severity, 'explanation': explanation})

    for orig, name, typ, unit, desc, role in FIELDS:
        s = d[name]
        blank = s.isna() | s.map(lambda x: isinstance(x, str) and not x.strip())
        if typ == 'number':
            d[name] = pd.to_numeric(s, errors='coerce')
            invalid = ~blank & (d[name].isna() | ~np.isfinite(d[name]))
            check('invalid_number:' + name, invalid, 'quarantine', 'Non-numeric/non-finite value in numeric field')
        elif typ == 'date':
            d[name] = pd.to_datetime(s, errors='coerce')
            check('invalid_date:' + name, ~blank & d[name].isna(), 'quarantine', 'Unparseable date')
        else:
            d[name] = s.map(lambda x: None if pd.isna(x) else str(x).strip())
        mandatory = name in ['agreement_id', 'agreement_date', 'seizure_date', 'sold_date', 'asset_cost_at_disbursal', 'sale_amount']
        check('missing:' + name, blank, 'quarantine' if mandatory else 'review', 'Missing source value')

    check('duplicate_agreement_id', d.agreement_id.duplicated(keep=False), 'quarantine', 'All members retained in quarantine; no arbitrary winner')
    check('agreement_after_seizure', d.agreement_date > d.seizure_date, 'quarantine', 'Invalid event chronology')
    check('seizure_after_sale', d.seizure_date > d.sold_date, 'quarantine', 'Invalid event chronology')
    for c in ['asset_cost_at_disbursal', 'sale_amount']:
        check('nonpositive:' + c, d[c] <= 0, 'quarantine', 'Valuation cost/target must be positive')
    for c in ['customer_net_salary_raw', 'loan_amount', 'outstanding_balance_at_liquidation', 'traffic_challan_amount', 'asset_age_months_at_seizure_raw', 'months_spent_in_yard_raw']:
        check('negative:' + c, d[c] < 0, 'review', 'Investigate recorded negative value; not auto-deleted')
    check('nonpositive_tenure', d.tenure_months <= 0, 'review', 'Cannot use for policy amortization')
    check('ltv_outside_zero_one', (d.ltv <= 0) | (d.ltv > 1), 'review', 'Confirm units/financed costs; not an automatic valuation exclusion')
    check('cibil_negative_special_code', d.customer_cibil_score_raw < 0, 'review', 'Unresolved negative source code; usable score set null')
    check('income_nonpositive', d.customer_net_salary_raw <= 0, 'review', 'Affordability unavailable without verified positive income and frequency')

    elapsed = (d.seizure_date - d.agreement_date).dt.total_seconds() / 86400 / 30
    yard_days = (d.sold_date - d.seizure_date).dt.total_seconds() / 86400
    ltv_calc = d.loan_amount / d.asset_cost_at_disbursal.replace(0, np.nan)
    age_err = (d.asset_age_months_at_seizure_raw - elapsed).abs()
    yard_err = (d.months_spent_in_yard_raw - yard_days / 30).abs()
    ltv_err = (d.ltv - ltv_calc).abs()
    check('age_elapsed_mismatch', age_err > 1e-8, 'review', 'Source age differs from 30-day agreement-to-seizure duration', 1e-8)
    check('yard_duration_mismatch', yard_err > 1e-8, 'review', 'Source yard months differs from seizure-to-sale days / 30', 1e-8)
    check('ltv_reconciliation_mismatch', ltv_err > 1e-6, 'review', 'LTV differs from loan amount / original cost', 1e-6)
    for dt, month, year in [('sold_date', 'sold_month_raw', 'sold_year_raw'), ('seizure_date', 'seizure_month_raw', 'seizure_year_raw')]:
        check('calendar_mismatch:' + month, d[month] != d[dt].dt.month, 'review', 'Redundant month differs from authoritative date')
        check('calendar_mismatch:' + year, d[year] != d[dt].dt.year, 'review', 'Redundant year differs from authoritative date')
    check('yard_over_six_months', d.months_spent_in_yard_raw > 6, 'review', 'Operational investigation threshold only; genuine long delays retained', 6)
    check('challan_exceeds_sale', d.traffic_challan_amount > d.sale_amount, 'review', 'Verify fine amount and recoverability; do not subtract automatically')
    check('liquidation_balance_exceeds_loan', d.outstanding_balance_at_liquidation > d.loan_amount, 'review', 'Could include arrears/charges; confirm balance definition')
    check('sale_exceeds_original_cost', d.sale_amount > d.asset_cost_at_disbursal, 'review', 'Verify valid outcome; do not cap target')

    d['customer_cibil_score_usable'] = d.customer_cibil_score_raw.where(d.customer_cibil_score_raw >= 0)
    d['cibil_special_code_flag'] = d.customer_cibil_score_raw.lt(0)
    d['customer_income_positive'] = d.customer_net_salary_raw.where(d.customer_net_salary_raw > 0)
    d['income_nonpositive_flag'] = d.customer_net_salary_raw.le(0)
    d['months_since_agreement_at_seizure'] = elapsed
    d['realized_yard_days'] = yard_days
    d['ltv_recomputed'] = ltv_calc
    quarantine_rows = {x['source_excel_row'] for x in issues if x['severity'] == 'quarantine'}
    qmask = d.source_excel_row.isin(quarantine_rows)
    clean = d.loc[~qmask].copy()
    quarantined = d.loc[qmask].copy()
    write_jsonl(out / 'data/canonical_records.jsonl', clean)
    write_jsonl(out / 'data/quarantined_records.jsonl', quarantined)
    raw_snapshot = raw.copy()
    raw_snapshot.insert(0, 'source_excel_row', excel_rows)
    write_jsonl(out / 'data/source_values.jsonl', raw_snapshot)
    write_json(out / 'reports/quality_checks.json', checks)
    issue_frame = pd.DataFrame(issues, columns=['source_excel_row','agreement_id','rule','severity','explanation'])
    write_jsonl(out / 'reports/row_issues.jsonl', issue_frame)

    profiles = []
    dictionary = []
    for pos, (orig, name, typ, unit, desc, role) in enumerate(FIELDS, 1):
        s = d[name]
        p = {'source': orig, 'canonical': name, 'source_excel_column': openpyxl.utils.get_column_letter(pos),
             'missing': int(s.isna().sum()), 'unique': int(s.nunique()), 'type': typ}
        if typ == 'number':
            p['statistics'] = {k: scalar(v) for k, v in s.describe(percentiles=[.01,.05,.25,.5,.75,.95,.99]).items()}
        elif typ == 'date':
            p['min'] = scalar(s.min()); p['max'] = scalar(s.max())
        elif role != 'identifier':
            p['frequencies'] = {str(k): int(v) for k,v in s.value_counts(dropna=False).items()}
        profiles.append(p)
        seizure, origination = availability(role)
        dictionary.append({'source_name': orig, 'canonical_name': name, 'type': typ, 'unit': unit,
                           'definition': desc, 'role': role, 'seizure_use': seizure, 'origination_use': origination,
                           'observed_codes': p.get('frequencies'), 'observed_codes_are_not_a_business_codebook': True,
                           'null_count_source': p['missing'], 'source_timestamp': 'not supplied per field',
                           'treatment': 'Normalize names/types only; preserve source codes. See derived fields for null-safe views.'})
    for name, typ, unit, desc, role in DERIVED:
        a,b = availability(role)
        dictionary.append({'source_name': None, 'canonical_name': name, 'type': typ, 'unit': unit,
                           'definition': desc, 'role': role, 'seizure_use': a, 'origination_use': b})
    write_json(out / 'reports/column_profiles.json', profiles)
    write_json(out / 'configs/data_dictionary.json', dictionary)

    base = [x[1] for x in FIELDS if x[5] == 'base_candidate'] + ['months_since_agreement_at_seizure']
    conditional = [x[1] for x in FIELDS if x[5] == 'conditional']
    forbidden = sorted(set(clean.columns) - set(base) - set(conditional))
    policy = {'status': 'Phase 1 candidate feature policy; not certified timestamp provenance',
              'seizure_core_candidate_features': base,
              'origination_core_candidate_features': [x for x in base if x != 'months_since_agreement_at_seizure'],
              'excluded_pending_definition_or_timestamp': conditional,
              'excluded_from_core_valuation': forbidden,
              'target': 'sale_amount', 'time_anchor': 'seizure_date', 'label_event_date': 'sold_date',
              'system_label_arrival_timestamp': None,
              'warning': 'Canonical records are an audit dataset, not a ready-made training matrix. Select explicit candidate fields only.'}
    write_json(out / 'configs/feature_policy.json', policy)

    monthly = d.groupby(d.sold_date.dt.strftime('%Y-%m')).agg(records=('agreement_id','size'), median_sale=('sale_amount','median'), median_yard_months=('months_spent_in_yard_raw','median'))
    monthly_records = [{'sold_month': month, **{k: scalar(v) for k,v in row.items()}} for month,row in monthly.iterrows()]
    segment_rows = []
    for c in ['asset_fuel_type', 'asset_model', 'customer_region', 'customer_state']:
        for key, g in d.groupby(c, dropna=False):
            segment_rows.append({'field': c, 'segment': scalar(key), 'records': len(g),
                                 'median_sale': scalar(g.sale_amount.median()), 'median_original_cost': scalar(g.asset_cost_at_disbursal.median()),
                                 'median_sale_to_cost_ratio': scalar((g.sale_amount/g.asset_cost_at_disbursal).median())})
    write_json(out / 'reports/monthly_profile.json', monthly_records)
    write_json(out / 'reports/segment_profile.json', segment_rows)
    condition_cols = ['asset_body_condition','asset_tyre_condition','asset_general_condition','asset_engine_condition']
    same_conditions = int(d[condition_cols].nunique(axis=1).eq(1).sum())
    gross_shortfall = (d.outstanding_balance_at_liquidation - d.sale_amount).clip(lower=0)
    summary = {
        'source_records': len(d), 'source_columns': len(FIELDS), 'canonical_columns': len(clean.columns),
        'retained_records': len(clean), 'quarantined_records': len(quarantined), 'unique_agreements': int(d.agreement_id.nunique()),
        'blank_source_cells': int(raw.isna().sum().sum()),
        'exact_duplicate_rows_excluding_id': int(raw.drop(columns=['Agmt Id']).duplicated(keep=False).sum()),
        'review_records': len({x['source_excel_row'] for x in issues if x['severity']=='review'}),
        'review_issue_occurrences': sum(x['severity']=='review' for x in issues),
        'cibil_negative_count': int(d.cibil_special_code_flag.sum()), 'income_nonpositive_count': int(d.income_nonpositive_flag.sum()),
        'fuel_counts': {str(k):int(v) for k,v in d.asset_fuel_type.value_counts().items()},
        'age_equals_agreement_elapsed_count': int(age_err.le(1e-8).sum()), 'age_max_absolute_difference': scalar(age_err.max()),
        'yard_equals_elapsed_count': int(yard_err.le(1e-8).sum()), 'yard_max_absolute_difference': scalar(yard_err.max()),
        'ltv_matches_count': int(ltv_err.le(1e-6).sum()), 'ltv_max_absolute_difference': scalar(ltv_err.max()),
        'all_four_condition_codes_equal_count': same_conditions,
        'sold_date_min': scalar(d.sold_date.min()), 'sold_date_max': scalar(d.sold_date.max()),
        'sale_amount_sum': scalar(d.sale_amount.sum()), 'sale_amount_median': scalar(d.sale_amount.median()),
        'gross_shortfall_sum': scalar(gross_shortfall.sum()), 'gross_shortfall_positive_records': int(gross_shortfall.gt(0).sum()),
        'gross_shortfall_definition': 'sum(max(liquidation balance - sale amount,0)); descriptive proxy, not realized economic loss/LGD',
        'checks_performed': len(checks), 'source_sha256_before': before,
    }
    write_json(out / 'reports/summary.json', summary)
    dictionary_md = ['# Phase 1 data dictionary', '', 'Source: Analytics Case Study Dataset.xlsx, Sheet1. All 41 source columns plus 8 traceability/derived columns. Observed codes are not verified definitions.', '',
                     '| Source field | Canonical field | Type / unit | Definition | Seizure use | Origination use |',
                     '|---|---|---|---|---|---|']
    for item in dictionary:
        dictionary_md.append('| ' + ' | '.join(str(item.get(k) or '(derived)').replace('|','/') for k in ['source_name','canonical_name']) + ' | ' + item['type'] + ' / ' + item['unit'] + ' | ' + item['definition'] + ' | ' + item['seizure_use'] + ' | ' + item['origination_use'] + ' |')
    (out/'reports/Data_Dictionary.md').write_text('\n'.join(dictionary_md)+'\n', encoding='utf-8')

    report = [
        '# Phase 1: dataset audit and preparation', '',
        'Source: Analytics Case Study Dataset.xlsx, Sheet1. This report describes the complete supplied file, not model performance. Source workbook unchanged.', '',
        '## Outcome', '',
        f'- {len(d):,} source rows, {len(FIELDS)} source columns, {d.agreement_id.nunique():,} unique agreement IDs.',
        f'- {len(clean):,} rows retained; {len(quarantined):,} quarantined; {summary["review_records"]:,} distinct rows have one or more review flags.',
        f'- {summary["blank_source_cells"]:,} blank source cells. Special codes and zero income still require semantic handling.',
        '- Canonical records are provided as JSON Lines, with one object per agreement and explicit numeric/null/boolean types. No rows are silently imputed, capped or deduplicated.',
        '- Modelling, chronological split assignment and model fitting belong to Phase 2 onward and have not been performed.', '',
        '## Important findings', '',
        f'1. **Age definition:** {summary["age_equals_agreement_elapsed_count"]:,}/{len(d):,} reported age values equal (seizure date - agreement date)/30 within 1e-8 months. Use `months_since_agreement_at_seizure` for this quantity. It is not confirmed manufacture age.',
        f'2. **Realized yard duration:** {summary["yard_equals_elapsed_count"]:,}/{len(d):,} yard values equal sale-minus-seizure days/30. This is unavailable at seizure and must stay out of the seizure model.',
        f'3. **CIBIL:** {summary["cibil_negative_count"]:,} records ({summary["cibil_negative_count"]/len(d):.1%}) have a negative special code. Preserve the raw code and use null in the derived usable view. Its business meaning remains unresolved.',
        f'4. **Income:** {summary["income_nonpositive_count"]:,} records contain zero/negative recorded income. These cannot establish affordability. Income frequency and existing obligations are missing.',
        f'5. **EV sample:** {summary["fuel_counts"].get("EV",0):,} EV records ({summary["fuel_counts"].get("EV",0)/len(d):.2%}); use pooled estimates initially and disclose small segment samples.',
        f'6. **LTV reconciliation:** {summary["ltv_matches_count"]:,}/{len(d):,} supplied values match loan/original-cost within 1e-6. Retain fraction units.',
        f'7. **Condition patterns:** all four codes agree in {same_conditions:,} records ({same_conditions/len(d):.1%}). Investigate coding/inspection practice; this alone does not establish fabricated data.',
        '8. **Population:** observed liquidation sales exclude the performing-loan population and unsold seizures. Default probability, full credit LGD and causal policy uplift cannot be identified from this file alone.',
        f'9. **Time coverage:** sale outcomes span {str(d.sold_date.min().date())} through {str(d.sold_date.max().date())}; long-horizon forecasts require additional evidence.', '',
        '## Quality checks with findings', '', '| Rule | Rows | Handling |', '|---|---:|---|']
    for c in checks:
        if c['affected_rows']:
            report.append(f'| {c["rule"]} | {c["affected_rows"]:,} | {c["severity"]}: {c["explanation"]} |')
    report += ['', 'The complete rule list, including zero-findings checks, is in quality_checks.json. Flags overlap; never add flagged counts to estimate unique affected agreements. Review flags do not automatically quarantine a vehicle.', '',
               '## Monthly sale coverage', '', '| Sale month | Records | Median sale (INR) | Median realized yard months |', '|---|---:|---:|---:|']
    for r in monthly_records:
        report.append(f'| {r["sold_month"]} | {r["records"]:,.0f} | {r["median_sale"]:,.0f} | {r["median_yard_months"]:.2f} |')
    report += ['', 'Monthly medians are descriptive and mix-dependent; they are not a quality-adjusted market price index.', '',
               '## Cleaning performed', '',
               '- Assigned stable canonical field names and retained source worksheet row references.',
               '- Preserved all source field values in a separate value snapshot. Dates are serialized as ISO timestamps; Excel display formatting is not part of that snapshot.',
               '- Converted canonical flags/categories to strings and numeric/date fields to typed values; stripped surrounding whitespace only.',
               '- Preserved raw CIBIL/income values; appended null-safe derived views and explicit flags without guessing meanings.',
               '- Derived agreement elapsed time, realized yard days and an independently recomputed LTV.',
               '- Stored field-level policy, row-level issues, full profiles and file hashes.',
               '- Performed no target-based imputation, category pooling, outlier deletion, train/test split or model fitting.', '',
               '## Questions and Phase 2 handoff', '',
               'Read Open_Definitions.md for the unresolved source-owner questions. Safe preparatory work can continue with the explicit feature whitelist; inspection/document flags stay excluded until timing is verified.',
               'Phase 2 must construct temporal splits with label availability and keep repeated aggregates/transforms inside training folds. Do not pass the entire canonical audit table to a model.',
               'The proposed core uses original cost, vehicle/model/fuel and location plus elapsed agreement time at seizure. Their immutable/historical status is still an explicit POC assumption, not verified per-field provenance.', '',
               '## Verification', '',
               'The companion verification script checks source hashes, row counts, source-to-canonical value preservation, derived calculations, issue reconciliation and forbidden feature exclusion. verification.json records the executed result. No prediction accuracy is asserted.']
    (out/'reports/Phase_1_Report.md').write_text('\n'.join(report)+'\n', encoding='utf-8')
    after = sha(source)
    if before != after:
        raise AssertionError('Source changed during read-only audit')
    package_files = [p for folder in ['data','reports','configs'] for p in (out/folder).rglob('*') if p.is_file() and p.name != 'verification.json']
    manifest = {'schema_version':'1.0','created_utc':datetime.now(timezone.utc).isoformat(),
                'source_file':source.name,'source_sheet':'Sheet1','source_sha256':before,'source_unchanged':before==after,
                'source_records':len(d),'retained_records':len(clean),'quarantined_records':len(quarantined),
                'python_version':platform.python_version(), 'packages':{x:importlib.metadata.version(x) for x in ['pandas','numpy','openpyxl']},
                'serialization':'JSON Lines UTF-8; ISO dates, numeric JSON values and explicit nulls',
                'source_snapshot_note':'Cell values only; source workbook formatting not reproduced',
                'audit_script_sha256':sha(Path(__file__)),
                'files':{p.relative_to(out).as_posix():sha(p) for p in sorted(package_files)}}
    write_json(out/'manifest.json',manifest)
    print(json.dumps(summary, indent=2, allow_nan=False))


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args=parser.parse_args()
    audit(args.source,args.output)
